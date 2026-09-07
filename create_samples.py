"""Script to generate realistic sample resumes (PDF and DOCX)."""

import os
from pathlib import Path
import docx
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

SAMPLES_DIR = Path(__file__).resolve().parent / "samples"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)


def create_resume_1_pdf():
    """Sample 1: Tech Fresher PDF."""
    path = SAMPLES_DIR / "resume_1.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter

    # Margins and layout
    y = height - 50

    # Header / Name
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, y, "Aarav Sharma")
    y -= 22

    # Contact Line
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Email: aarav.sharma@example.com | Phone: +91 98765 43210")
    y -= 15
    c.drawString(50, y, "LinkedIn: https://linkedin.com/in/aaravsharma | GitHub: https://github.com/aaravsharma")
    y -= 25

    # Horizontal divider
    c.setLineWidth(1)
    c.line(50, y, width - 50, y)
    y -= 20

    # Education Section
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "EDUCATION")
    y -= 18

    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "B.Tech in Computer Science")
    c.setFont("Helvetica", 11)
    c.drawRightString(width - 50, y, "2020 - 2024")
    y -= 15

    c.drawString(50, y, "ABC Institute of Technology")
    y -= 25

    # Experience Section
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "WORK EXPERIENCE")
    y -= 18

    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Software Engineering Intern")
    c.setFont("Helvetica", 11)
    c.drawRightString(width - 50, y, "Jan 2024 - Jun 2024")
    y -= 15

    c.drawString(50, y, "Nexus Tech Solutions")
    y -= 15

    c.setFont("Helvetica", 10)
    c.drawString(60, y, "- Developed REST APIs using Python and FastAPI framework.")
    y -= 14
    c.drawString(60, y, "- Containerized microservices using Docker and integrated PostgreSQL database.")
    y -= 25

    # Skills Section
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "TECHNICAL SKILLS")
    y -= 18

    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Programming Languages: Python, Java, SQL, C++")
    y -= 15
    c.drawString(50, y, "Web Frameworks: FastAPI, React, HTML, CSS")
    y -= 15
    c.drawString(50, y, "Databases & Tools: PostgreSQL, Docker, Git, Linux")
    y -= 25

    # Projects Section
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "PROJECTS")
    y -= 18

    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Automated Document Analyzer")
    y -= 14
    c.setFont("Helvetica", 10)
    c.drawString(60, y, "- Built an offline information extraction tool using Python and regex.")

    c.save()
    print(f"Generated {path}")


def create_resume_2_docx():
    """Sample 2: Senior Software Engineer DOCX."""
    path = SAMPLES_DIR / "resume_2.docx"
    doc = Document()

    # Header
    p_name = doc.add_paragraph()
    run_name = p_name.add_run("Sarah Jenkins")
    run_name.bold = True
    run_name.font.size = docx.shared.Pt(20)

    p_contact = doc.add_paragraph("sarah.jenkins@techmail.org | (415) 555-0199 | San Francisco, CA")
    doc.add_paragraph("LinkedIn: https://www.linkedin.com/in/sarahjenkins-dev | GitHub: https://github.com/sarahjenkins")

    # Experience Section
    h_exp = doc.add_heading(level=1)
    h_exp.add_run("Work Experience").bold = True

    p1 = doc.add_paragraph()
    p1.add_run("Full Stack Developer").bold = True
    p1.add_run(" - Apex Cloud Systems (Aug 2022 - Present)")
    doc.add_paragraph("• Architected scalable web applications using TypeScript, React, and Node.js.")
    doc.add_paragraph("• Deployed cloud services on AWS with Docker and Kubernetes.")

    p2 = doc.add_paragraph()
    p2.add_run("Software Engineer").bold = True
    p2.add_run(" - Pinnacle Software Labs (Jul 2020 - Jul 2022)")
    doc.add_paragraph("• Built high-performance microservices using Python, Redis, and MongoDB.")
    doc.add_paragraph("• Configured CI/CD pipelines with GitHub Actions.")

    # Education Section
    h_edu = doc.add_heading(level=1)
    h_edu.add_run("Education").bold = True

    p_edu1 = doc.add_paragraph()
    p_edu1.add_run("Master of Science in Computer Science").bold = True
    p_edu1.add_run(" | Stanford University (2020)")

    p_edu2 = doc.add_paragraph()
    p_edu2.add_run("Bachelor of Science").bold = True
    p_edu2.add_run(" | University of California (2018)")

    # Skills Section
    h_skills = doc.add_heading(level=1)
    h_skills.add_run("Technical Skills").bold = True

    doc.add_paragraph("Languages: JavaScript, TypeScript, Python, SQL")
    doc.add_paragraph("Frameworks & Tools: React, Node.js, Express, AWS, Docker, Kubernetes, Redis, MongoDB, GraphQL, Git")

    doc.save(str(path))
    print(f"Generated {path}")


def create_resume_3_pdf():
    """Sample 3: Minimal Data Science Resume with missing optional fields."""
    path = SAMPLES_DIR / "resume_3.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter

    y = height - 50

    # Header / Name
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, y, "Vikram Patel")
    y -= 20

    # Contact (No LinkedIn, No GitHub)
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Email: vikram.patel99@gmail.com | Phone: 9823456789")
    y -= 25

    c.setLineWidth(1)
    c.line(50, y, width - 50, y)
    y -= 25

    # Education
    c.setFont("Helvetica-Bold", 13)
    c.drawString(50, y, "EDUCATION")
    y -= 18

    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "B.Sc in Statistics")
    c.setFont("Helvetica", 10)
    c.drawRightString(width - 50, y, "2023")
    y -= 14

    c.drawString(50, y, "City College")
    y -= 30

    # Skills
    c.setFont("Helvetica-Bold", 13)
    c.drawString(50, y, "TECHNICAL SKILLS")
    y -= 18

    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Python, SQL, Pandas, NumPy, Machine Learning, Data Analysis, Tableau")
    y -= 30

    # Summary
    c.setFont("Helvetica-Bold", 13)
    c.drawString(50, y, "ACADEMIC SUMMARY")
    y -= 18
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Recent graduate with strong foundation in statistical modeling and data science.")

    c.save()
    print(f"Generated {path}")


if __name__ == "__main__":
    create_resume_1_pdf()
    create_resume_2_docx()
    create_resume_3_pdf()
