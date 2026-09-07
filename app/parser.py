"""Document parsing module for PDF and DOCX files with visual layout reconstruction."""

import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import docx
import pypdf

try:
    import pymupdf  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz as pymupdf  # type: ignore
        HAS_PYMUPDF = True
    except ImportError:
        HAS_PYMUPDF = False


class DocumentParserError(Exception):
    """Base exception for document parsing errors."""
    pass


class UnsupportedFileTypeError(DocumentParserError):
    """Raised when an uploaded file is not PDF or DOCX."""
    pass


class EmptyDocumentError(DocumentParserError):
    """Raised when the uploaded file is empty (0 bytes)."""
    pass


class CorruptedFileError(DocumentParserError):
    """Raised when a file cannot be parsed due to corruption."""
    pass


class NoTextFoundError(DocumentParserError):
    """Raised when a file contains no readable digital text (e.g. scanned image)."""
    pass


def extract_pdf_blocks(file_stream: io.BytesIO) -> Tuple[List[Dict[str, Any]], Dict[int, Tuple[float, float]]]:
    """Extracts structured text spans with bounding boxes and typography from PDF.
    
    Args:
        file_stream: In-memory byte stream of the PDF file.
        
    Returns:
        Tuple of (list of text spans with coordinates/styling, map of page dimensions {page_idx: (width, height)}).
        
    Raises:
        CorruptedFileError: If the PDF structure is invalid or password-protected.
        NoTextFoundError: If no readable digital text was found.
    """
    if not HAS_PYMUPDF:
        raise CorruptedFileError("PyMuPDF library is required for PDF block extraction.")

    file_stream.seek(0)
    try:
        doc = pymupdf.open(stream=file_stream.read(), filetype="pdf")
    except Exception as exc:
        raise CorruptedFileError(f"Failed to open PDF document: {str(exc)}") from exc

    if doc.is_encrypted:
        try:
            doc.authenticate("")
        except Exception:
            raise CorruptedFileError("PDF is password-protected and cannot be processed.")

    all_spans: List[Dict[str, Any]] = []
    page_dims: Dict[int, Tuple[float, float]] = {}

    try:
        for page_idx, page in enumerate(doc):
            rect = page.rect
            page_dims[page_idx] = (rect.width, rect.height)
            d = page.get_text("dict")
            for block in d.get("blocks", []):
                if "lines" in block:
                    for line in block["lines"]:
                        for span in line.get("spans", []):
                            text = span.get("text", "").strip()
                            if text:
                                bbox = span.get("bbox", (0.0, 0.0, 0.0, 0.0))
                                font_flags = span.get("flags", 0)
                                font_name = span.get("font", "")
                                is_bold = bool("bold" in font_name.lower() or (font_flags & 2 != 0))
                                all_spans.append({
                                    "text": text,
                                    "bbox": bbox,
                                    "size": float(span.get("size", 10.0)),
                                    "bold": is_bold,
                                    "flags": font_flags,
                                    "font": font_name,
                                    "page": page_idx
                                })
    except Exception as exc:
        raise CorruptedFileError(f"Error extracting PDF blocks: {str(exc)}") from exc

    if not all_spans:
        raise NoTextFoundError("No readable text found in PDF. The document may be empty or a scanned image.")

    return all_spans, page_dims


def _cluster_spans_into_lines(spans: List[Dict[str, Any]], y_tolerance: float = 4.5) -> List[str]:
    """Clusters spans with similar vertical baselines into coherent visual lines."""
    if not spans:
        return []

    # Sort spans primarily by approximate y-baseline, then by x0
    sorted_spans = sorted(
        spans,
        key=lambda s: (round((s["bbox"][1] + s["bbox"][3]) / (2.0 * y_tolerance)) * y_tolerance, s["bbox"][0])
    )

    lines: List[str] = []
    curr_line: List[Dict[str, Any]] = []
    curr_y: Optional[float] = None

    for s in sorted_spans:
        mid_y = (s["bbox"][1] + s["bbox"][3]) / 2.0
        if curr_y is None or abs(mid_y - curr_y) > y_tolerance:
            if curr_line:
                curr_line.sort(key=lambda item: item["bbox"][0])
                line_str = curr_line[0]["text"]
                for i in range(1, len(curr_line)):
                    gap = curr_line[i]["bbox"][0] - curr_line[i - 1]["bbox"][2]
                    # If there is a wide horizontal gap (e.g. institution on left, date/grade on right)
                    if gap > 22.0:
                        line_str += "  |  " + curr_line[i]["text"]
                    else:
                        line_str += " " + curr_line[i]["text"]
                lines.append(line_str.strip())
            curr_line = [s]
            curr_y = mid_y
        else:
            curr_line.append(s)
            curr_y = (curr_y * len(curr_line) + mid_y) / (len(curr_line) + 1)

    if curr_line:
        curr_line.sort(key=lambda item: item["bbox"][0])
        line_str = curr_line[0]["text"]
        for i in range(1, len(curr_line)):
            gap = curr_line[i]["bbox"][0] - curr_line[i - 1]["bbox"][2]
            if gap > 22.0:
                line_str += "  |  " + curr_line[i]["text"]
            else:
                line_str += " " + curr_line[i]["text"]
        lines.append(line_str.strip())

    return lines


def reconstruct_reading_order(
    spans: List[Dict[str, Any]],
    page_dims: Optional[Dict[int, Tuple[float, float]]] = None
) -> str:
    """Reconstructs reading order from PDF spans using coordinate layout analysis.
    
    Supports both single-column and multi-column visual layouts.
    
    Args:
        spans: List of text span dicts containing 'text', 'bbox', 'size', 'page'.
        page_dims: Optional mapping of page index to (width, height).
        
    Returns:
        Structured, ordered resume text.
    """
    if not spans:
        return ""

    # Group spans by page
    pages_spans: Dict[int, List[Dict[str, Any]]] = {}
    for span in spans:
        p = span.get("page", 0)
        pages_spans.setdefault(p, []).append(span)

    reconstructed_page_texts: List[str] = []

    for page_idx in sorted(pages_spans.keys()):
        p_spans = pages_spans[page_idx]
        width, height = (612.0, 792.0)
        if page_dims and page_idx in page_dims:
            width, height = page_dims[page_idx]

        # Check for two-column layout:
        is_two_column = False
        best_split_x = None

        if len(p_spans) > 20:
            for split_ratio in (0.30, 0.33, 0.35, 0.40, 0.50):
                test_split = width * split_ratio
                left_count = 0
                right_count = 0
                crossing_count = 0
                body_spans = [s for s in p_spans if 0.15 * height <= s["bbox"][1] <= 0.95 * height]
                for s in body_spans:
                    x0, _, x1, _ = s["bbox"]
                    if x1 <= test_split + 5.0:
                        left_count += 1
                    elif x0 >= test_split - 5.0:
                        right_count += 1
                    else:
                        crossing_count += 1

                total_body = len(body_spans)
                if total_body >= 15 and left_count >= 5 and right_count >= 8 and crossing_count <= 2:
                    is_two_column = True
                    best_split_x = test_split
                    break

        if is_two_column and best_split_x is not None:
            header_threshold = 0.18 * height
            header_spans = [s for s in p_spans if s["bbox"][3] <= header_threshold]
            body_left = [s for s in p_spans if s["bbox"][3] > header_threshold and s["bbox"][2] <= best_split_x + 5.0]
            body_right = [s for s in p_spans if s["bbox"][3] > header_threshold and s["bbox"][0] >= best_split_x - 5.0]
            handled_ids = {id(s) for s in header_spans + body_left + body_right}
            others = [s for s in p_spans if id(s) not in handled_ids]

            page_lines = []
            if header_spans:
                page_lines.extend(_cluster_spans_into_lines(header_spans))
            if body_left:
                page_lines.extend(_cluster_spans_into_lines(body_left))
            if body_right:
                page_lines.extend(_cluster_spans_into_lines(body_right))
            if others:
                page_lines.extend(_cluster_spans_into_lines(others))
            reconstructed_page_texts.append("\n".join(page_lines))
        else:
            # Single-column or standard layout
            page_lines = _cluster_spans_into_lines(p_spans)
            reconstructed_page_texts.append("\n".join(page_lines))

    return "\n\n".join(reconstructed_page_texts).strip()


def parse_pdf(file_stream: io.BytesIO) -> str:
    """Extracts text from a PDF byte stream with layout reconstruction.
    
    Args:
        file_stream: In-memory byte stream of the PDF file.
        
    Returns:
        Extracted, cleanly ordered text as a single string.
        
    Raises:
        CorruptedFileError: If PDF structure is invalid.
        NoTextFoundError: If no digital text could be extracted.
    """
    if HAS_PYMUPDF:
        try:
            spans, dims = extract_pdf_blocks(file_stream)
            ordered_text = reconstruct_reading_order(spans, dims)
            if ordered_text.strip():
                return ordered_text
        except (CorruptedFileError, NoTextFoundError):
            raise
        except Exception:
            pass

    # Fallback to pypdf if PyMuPDF fails or is unavailable
    file_stream.seek(0)
    try:
        reader = pypdf.PdfReader(file_stream)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise CorruptedFileError("PDF is password-protected and cannot be processed.")

        extracted_pages = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text.strip())

        full_text = "\n\n".join(extracted_pages).strip()
        if not full_text:
            raise NoTextFoundError(
                "No readable text found in PDF. The document may be empty or a scanned image."
            )
        return full_text
    except (CorruptedFileError, NoTextFoundError):
        raise
    except Exception as exc:
        raise CorruptedFileError(f"Failed to parse PDF document: {str(exc)}") from exc


def parse_docx(file_stream: io.BytesIO) -> str:
    """Extracts text from a DOCX byte stream using python-docx.
    
    Inspects both standard paragraphs and table cell contents.
    
    Args:
        file_stream: In-memory byte stream of the DOCX file.
        
    Returns:
        Extracted text as a single string.
        
    Raises:
        CorruptedFileError: If the DOCX structure is invalid.
        NoTextFoundError: If no text was found in paragraphs or tables.
    """
    try:
        doc = docx.Document(file_stream)
        lines = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        # Extract text from tables (common in CV templates)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate repeated cell text from merged cells
                unique_cells = []
                for cell in row_cells:
                    if not unique_cells or cell != unique_cells[-1]:
                        unique_cells.append(cell)
                if unique_cells:
                    lines.append(" | ".join(unique_cells))

        full_text = "\n".join(lines).strip()
        if not full_text:
            raise NoTextFoundError(
                "No readable text found in DOCX file. The document is empty."
            )
        return full_text
    except (CorruptedFileError, NoTextFoundError):
        raise
    except Exception as exc:
        raise CorruptedFileError(f"Failed to parse DOCX document: {str(exc)}") from exc


def extract_document_layout(file_bytes: bytes, filename: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Extracts both reconstructed text and structured block metadata from file.
    
    Args:
        file_bytes: Raw bytes of the uploaded document.
        filename: Document filename.
        
    Returns:
        Tuple of (ordered_text, list_of_spans_or_blocks).
    """
    if not file_bytes or len(file_bytes.strip()) == 0:
        raise EmptyDocumentError("The uploaded file is empty (0 bytes).")

    ext = Path(filename).suffix.lower()
    stream = io.BytesIO(file_bytes)

    if ext == ".pdf":
        if HAS_PYMUPDF:
            try:
                spans, dims = extract_pdf_blocks(stream)
                text = reconstruct_reading_order(spans, dims)
                return text, spans
            except Exception:
                pass
        text = parse_pdf(stream)
        return text, []
    elif ext == ".docx":
        return parse_docx(stream), []
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word (.docx) documents are accepted."
        )


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extracts text from PDF or DOCX file content.
    
    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename: Original filename (used to inspect extension).
        
    Returns:
        Extracted and trimmed raw text string.
        
    Raises:
        EmptyDocumentError: If file_bytes is empty.
        UnsupportedFileTypeError: If file extension is not .pdf or .docx.
        CorruptedFileError: If file is damaged or unreadable.
        NoTextFoundError: If document contains no readable text.
    """
    text, _ = extract_document_layout(file_bytes, filename)
    return text

